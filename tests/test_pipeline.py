from pathlib import Path

import pytest

from crawl import export, ingest, layout as layout_mod, lint, timing as timing_mod
from crawl.fusion import emit
from crawl.look import Look
from crawl.model import Header, Legal, ListBlock, LogoRow, RolePair
from crawl.project import Build, Project, prepare, render
from crawl.timing import Timing, frames_to_tc, tc_to_frames

SAMPLE = Path(__file__).parent.parent / "examples" / "sample_crawl.csv"


# ---- ingest ----------------------------------------------------------------

def rows(*r):
    return [tuple(x) for x in r]


def test_parse_basic_shapes():
    c = ingest.parse(rows(
        ("", "CAST", ""),
        ("Lead", "", "A Person\nB Person"),
        ("", "", ""),
        ("", "Song Title", ""),
        ("Written by", "", "C Person"),
        ("", "", ""),
        *[(f"Partner {i}", "", "") for i in range(8)],
        ("", "", ""),
        ("[Logo: Studio]", "", ""),
        ("© 2026 Studio. All rights reserved.", "", ""),
    ))
    kinds = [type(b).__name__ for b in c.blocks()]
    assert kinds == ["Header", "RolePair", "Header", "RolePair", "ListBlock", "LogoRow", "Legal"]
    h1, rp, song = list(c.blocks())[:3]
    assert h1.level == 1 and song.level == 2
    assert rp.names == ["A Person", "B Person"]
    assert len(c.groups) == 4
    assert list(c.blocks())[4].names[0] == "Partner 0"


def test_logo_rows():
    assert ingest.logo_labels("[Logo: A] | [B LOGO] | Logo: C") == ["A", "B LOGO", "C"]
    assert ingest.logo_labels("Best Boy") == []
    c = ingest.parse(rows(("[Logo: A] | [Logo: B]", "", "")))
    (lr,) = list(c.blocks())
    assert isinstance(lr, LogoRow) and [l.label for l in lr.logos] == ["A", "B"]
    look = Look()
    L = layout_mod.layout(c, look, 1.0)
    assert len(L.logos) == 2 and L.logos[0].x + L.logos[0].w + look.logo_gap == L.logos[1].x
    assert abs((L.logos[0].x + L.logos[1].x + L.logos[1].w) / 2 - look.center_x) <= 1


def test_short_b_only_run_is_missing_names():
    c = ingest.parse(rows(("Head of Production", "", ""), ("Post Supervisor", "", "")))
    blocks = list(c.blocks())
    assert all(isinstance(b, RolePair) and not b.names for b in blocks)


def test_sample_counts():
    c = ingest.load(str(SAMPLE))
    k = c.counts()
    assert k["headers"] == 13
    assert k["list_names"] == 97 + 258
    assert k["logos"] == 2 and k["legal"] == 2


# ---- lint ------------------------------------------------------------------

def test_lint_finds_known_issues():
    c = ingest.load(str(SAMPLE))
    f = lint.lint(c)
    msgs = "\n".join(str(x) for x in f)
    assert "'Name 3'" in msgs and "(clear field)" not in msgs
    assert "'Head of Production' has no names" in msgs
    assert "'Jim Thelen' already listed" in msgs
    assert "'Lelsey Bohm' looks like 'Lesley Bohm Photography'" in msgs
    assert "First Assistant Director" in msgs
    assert "logo slot needs artwork" in msgs


# ---- timing ----------------------------------------------------------------

def test_tc_roundtrip():
    assert tc_to_frames("3:30", 23.976) == 5035
    assert tc_to_frames("210s", 23.976) == 5035
    assert tc_to_frames("5035f", 23.976) == 5035
    assert frames_to_tc(5035, 23.976).startswith("3:30")


def test_solve_runtime_integer_speed_exact_height():
    t = Timing(fps=23.976, runtime_frames=5035)
    s = timing_mod.solve(t, fixed_px=200, flex_px=48000, frame_height=2160)
    assert isinstance(s.speed, int)
    assert s.speed * s.travel_frames == s.content_height + 2160
    assert 0.9 < s.flex < 1.1


def test_solve_manual_rounds_up_to_whole_frames():
    t = Timing(mode="manual", speed=10, flex=1.0)
    s = timing_mod.solve(t, fixed_px=100, flex_px=48095, frame_height=2160)
    assert s.speed == 10
    assert (s.content_height + 2160) % 10 == 0


def test_solve_clamps_and_warns():
    # Long runtime, little content: even 1 px/frame needs the spacing stretched
    # past flex_max, so the solver clamps and says so.
    t = Timing(runtime_frames=5035)
    s = timing_mod.solve(t, fixed_px=200, flex_px=1000, frame_height=2160)
    assert s.speed == 1
    assert s.warnings and s.flex == t.flex_max


# ---- layout ----------------------------------------------------------------

def test_layout_is_linear_in_flex():
    c = ingest.load(str(SAMPLE))
    look = Look()
    fixed, flex = layout_mod.measure(c, look)
    for k in (0.8, 1.0, 1.3):
        h = layout_mod.layout(c, look, k).content_height
        assert abs(h - (fixed + k * flex)) <= 3


def test_panels_break_between_groups_and_are_even():
    c = ingest.load(str(SAMPLE))
    look = Look()
    L = layout_mod.layout(c, look, 1.0)
    L.target_height = L.content_height
    panels = layout_mod.split_panels(L, 12000)
    assert panels[0][0] == 0 and panels[-1][1] == L.target_height
    for top, bottom in panels:
        assert top % 2 == 0 and (bottom - top) % 2 == 0
        assert bottom - top <= 12000
    # No line straddles a cut.
    for _, cut in panels[:-1]:
        for ln in L.lines:
            assert not (cut - 60 < ln.y < cut + 60)


# ---- end to end ------------------------------------------------------------

def test_prepare_solves_sample_runtime():
    p = Project(source=str(SAMPLE), max_panel=12000)
    p.timing.runtime_frames = tc_to_frames("3:30", p.timing.fps)
    b = prepare(p)
    assert b.solution.total_frames == 5035
    assert b.layout.target_height + 2160 == b.solution.speed * b.solution.travel_frames + (b.layout.target_height % 2)


@pytest.mark.skipif(export.find_ffmpeg() is None, reason="ffmpeg not installed")
def test_render_movie_and_fusion(tmp_path):
    import subprocess
    p = Project(source=str(SAMPLE), output_dir=str(tmp_path), max_panel=12000,
                codec="h264", write_fusion=True)
    p.timing.mode = "manual"
    p.timing.speed = 200            # short: ~200 frames
    p.timing.head_hold = 5
    p.timing.tail_hold = 3
    b = prepare(p)
    paths = render(p, b)
    names = sorted(x.name for x in paths)
    assert "crawl.mp4" in names and "crawl.comp" in names and "crawl.setting" in names
    assert len([x for x in paths if x.suffix == ".png"]) == len(b.layout.panels)
    probe = subprocess.run(["ffprobe", "-v", "error", "-select_streams", "v", "-count_frames",
                            "-show_entries", "stream=nb_read_frames,width,height,r_frame_rate",
                            "-of", "csv=p=0", str(tmp_path / "crawl.mp4")],
                           capture_output=True, text=True).stdout.strip().split(",")
    assert probe[:3] == ["3840", "2160", "24000/1001"]
    assert int(probe[3]) == b.solution.total_frames
    comp = (tmp_path / "crawl.comp").read_text()
    assert "Point(0.5, ((time - 5) * 200" in comp


def test_export_frames_are_exact_crops():
    """Frame k of the movie must equal content rows [k*speed - 2160, k*speed)."""
    p = Project(source=str(SAMPLE))
    p.timing.mode = "manual"
    p.timing.speed = 300
    b = prepare(p)
    strip = export.render_strip(b.layout, p.look, colour=False)
    fh = p.look.height
    assert strip.shape == (fh + b.layout.target_height + fh, p.look.width)
    # Frame 0: content top just below the frame → all black.
    assert strip[0:fh].max() == 0
    # Frame 8: rows 8*300-2160 .. 8*300 of the content = rows 240..2400 of the strip after the pad.
    k = 8
    frame = strip[k * 300:k * 300 + fh]
    import numpy as np
    from PIL import Image
    from crawl.render import render_panel
    direct = np.asarray(render_panel(b.layout, p.look, k * 300 - fh, k * 300, 1.0, 1, colour=False))
    assert frame.max() > 0
    assert np.array_equal(frame, direct)


def test_clear_field_markers_are_empty_and_preserve_row_numbers():
    c = ingest.parse([
        ("", "(clear field)", ""),
        ("Director", "", "Alice"),
        ("( CLEAR FIELD )", "", "Bob"),
        ("Producer", "", "(clear fields)"),
    ])
    blocks = list(c.blocks())
    assert len(blocks) == 3
    assert blocks[0].row == 2 and blocks[0].names == ["Alice"]
    assert blocks[1].role == "" and blocks[1].names == ["Bob"]
    assert blocks[2].names == []


def test_xlsx_uses_only_crawl_and_rejects_missing_tab(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Opening Titles"
    wb.active.append(["", "Wrong sheet", "", "Wrong name"])
    ws = wb.create_sheet("Crawl")
    ws.append(["", "", "(clear field)", ""])
    ws.append(["", "Director", "", "Alice"])
    path = tmp_path / "credits.xlsx"
    wb.save(path)
    assert ingest.read_rows(str(path)) == [("", "", ""), ("Director", "", "Alice")]
    del wb["Crawl"]
    wb.save(path)
    with pytest.raises(ValueError, match="Worksheet 'Crawl' not found"):
        ingest.read_rows(str(path))
    wb.close()


def test_xlsx_excludes_hidden_template_rows_without_shifting_numbers(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Crawl'
    for row in [['', '', 'CAST', ''], ['', '', '', 'Name 3'],
                ['', '', 'Cast', ''], ['', 'Emma Lockhart', '', 'Name'],
                ['', 'Jamie Jaymes', '', 'Julie Benz'], ['', 'Visible role', '', 'Name 4']]:
        ws.append(row)
    for number in (2, 3, 4):
        ws.row_dimensions[number].hidden = True
    path = tmp_path / 'hidden.xlsx'
    wb.save(path)
    wb.close()
    rows = ingest.read_rows(str(path))
    assert len(rows) == 6
    assert rows[1:4] == [('', '', '')] * 3
    crawl = ingest.parse(rows)
    blocks = list(crawl.blocks())
    assert [(b.row, getattr(b, 'names', None)) for b in blocks] == [(1, None), (5, ['Julie Benz']), (6, ['Name 4'])]
    assert any(f.row == 6 and 'Name 4' in f.message for f in lint.lint(crawl))


# ---- layout overrides ------------------------------------------------------

def test_layout_kinds_override_b_only_rows():
    sheet = rows(("", "FOR STUDIO", ""), *[(f"Exec {i}", "", "") for i in range(8)])
    auto = ingest.parse(sheet)
    assert isinstance(list(auto.blocks())[1], ListBlock)

    roles = ingest.parse(sheet, kinds={i: "roles" for i in range(1, 9)})
    blocks = list(roles.blocks())[1:]
    assert all(isinstance(b, RolePair) and b.names == [] and b.names_pending for b in blocks)
    assert not any(f.check == "empty" for f in lint.lint(roles))
    L = layout_mod.layout(roles, Look())
    role_lines = [l for l in L.lines if l.style == "role"]
    assert len(role_lines) == 8 and all(l.anchor == "rs" for l in role_lines)

    # A short run forced to list, and a run split by a different kind.
    short = ingest.parse(rows(("A", "", ""), ("B", "", ""), ("C", "", "")), kinds={0: "list", 1: "list", 2: "list"})
    (lb,) = list(short.blocks())
    assert isinstance(lb, ListBlock) and lb.names == ["A", "B", "C"] and lb.columns == 3
    centred = ingest.parse(sheet, kinds={i: "center" for i in range(1, 5)})
    kinds = [type(b).__name__ for b in centred.blocks()]
    assert kinds == ["Header", "ListBlock", "RolePair", "RolePair", "RolePair", "RolePair"]
    assert list(centred.blocks())[1].columns == 1


# ---- logo placement --------------------------------------------------------

def _logo_png(path, colour, size=(200, 100)):
    from PIL import Image
    Image.new("RGBA", size, colour).save(path)
    return str(path)


def test_logo_settings_scale_offset_invert(tmp_path):
    import numpy as np
    from crawl.project import crawl_from_rows
    from crawl.render import render_panel
    dark = _logo_png(tmp_path / "dark.png", (20, 20, 20, 255))
    p = Project(rows=[["[Logo: A] | [Logo: B]", "", ""]], logos={"A": dark, "B": dark})
    look = p.look
    base = layout_mod.layout(crawl_from_rows(p, p.rows), look)
    a0, b0 = base.logos
    assert (a0.h, b0.h) == (look.logo_height, look.logo_height) and not a0.invert

    p.logo_settings = {"A": {"invert": True, "dx": 40, "dy": -30}, "B": {"scale": 0.5}}
    L = layout_mod.layout(crawl_from_rows(p, p.rows), look)
    a, b = L.logos
    assert a.invert and not b.invert
    assert b.h == look.logo_height // 2 and b.w == b0.w // 2
    # B is half height and centres on the row; A keeps its size and takes the offset.
    assert b.y == a0.y + (look.logo_height - b.h) // 2
    assert a.y == a0.y - 30
    # The row is narrower now, so its automatic start moved; the offset is on top of that.
    auto_x = round(look.center_x - (a.w + look.logo_gap + b.w) / 2)
    assert a.x == auto_x + 40

    img = np.asarray(render_panel(L, look, min(a.y, 0), a.y + a.h + 10, 1.0, 1, colour=False))
    top = min(a.y, 0)
    assert img[a.y - top + a.h // 2, a.x + a.w // 2] > 200      # inverted: dark → light
    assert img[b.y - top + b.h // 2, b.x + b.w // 2] < 60       # untouched: still dark


def test_coloured_logo_forces_colour_export(tmp_path):
    from crawl.project import crawl_from_rows
    grey = _logo_png(tmp_path / "grey.png", (200, 200, 200, 255))
    red = _logo_png(tmp_path / "red.png", (220, 30, 30, 255))
    p = Project(rows=[["[Logo: A]", "", ""]], logos={"A": grey})
    L = layout_mod.layout(crawl_from_rows(p, p.rows), p.look)
    assert not export.needs_colour(p.look, L)
    p.logos["A"] = red
    L = layout_mod.layout(crawl_from_rows(p, p.rows), p.look)
    assert export.needs_colour(p.look, L)


# ---- even logos ------------------------------------------------------------

def _art(path, canvas, ink, colour=(255, 255, 255, 255), bg=(0, 0, 0, 0)):
    from PIL import Image, ImageDraw
    im = Image.new("RGBA", canvas, bg)
    ImageDraw.Draw(im).rectangle([ink[0], ink[1], ink[2] - 1, ink[3] - 1], fill=colour)
    im.save(path)
    return str(path)


def test_ink_box_ignores_padding_and_resolution(tmp_path):
    from crawl import logos
    tight = _art(tmp_path / "tight.png", (400, 200), (0, 0, 400, 200))
    padded = _art(tmp_path / "padded.png", (1000, 1000), (100, 400, 900, 800))       # same 2:1 art, lots of margin
    black = _art(tmp_path / "black.png", (600, 400), (100, 100, 500, 300), bg=(0, 0, 0, 255))   # opaque, black margin
    boxed = _art(tmp_path / "boxed.png", (600, 400), (100, 100, 500, 300), bg=(60, 60, 60, 255))  # opaque grey box
    assert logos.ink_box(padded) == (100, 400, 900, 800)
    assert logos.ink_box(black) == (100, 100, 500, 300)
    assert logos.ink_box(boxed) == (0, 0, 600, 400)          # the box is the artwork
    assert logos.ink_box(str(tmp_path / "missing.png")) is None
    # Same artwork, different canvas and resolution → same size on the crawl.
    assert logos.size(tight, 300, 1.0) == logos.size(padded, 300, 1.0) == (600, 300)


def test_balance_evens_area_between_shapes(tmp_path):
    from crawl import logos
    wide = _art(tmp_path / "wide.png", (800, 100), (0, 0, 800, 100))       # 8:1 wordmark
    square = _art(tmp_path / "square.png", (64, 64), (0, 0, 64, 64))       # tiny 1:1 emblem
    # Equal heights: the wordmark has 8x the area.
    (ww, wh), (sw, sh) = logos.size(wide, 300, 0.0), logos.size(square, 300, 0.0)
    assert wh == sh == 300 and ww * wh == 8 * sw * sh
    # Balanced: areas are within the clamp of each other, wordmark shorter, emblem taller.
    (ww, wh), (sw, sh) = logos.size(wide, 300, 1.0), logos.size(square, 300, 1.0)
    assert wh < 300 < sh
    assert 0.5 <= (ww * wh) / (sw * sh) <= 2.0
    assert logos.MIN_FACTOR * 300 <= wh and sh <= logos.MAX_FACTOR * 300
    # The dial reaches the layout through Look.
    from crawl.project import crawl_from_rows
    p = Project(rows=[["[Logo: W] | [Logo: S]", "", ""]], logos={"W": wide, "S": square})
    p.look.logo_balance = 1.0
    w_item, s_item = layout_mod.layout(crawl_from_rows(p, p.rows), p.look).logos
    assert (w_item.w, w_item.h) == (ww, wh) and (s_item.w, s_item.h) == (sw, sh)
    assert w_item.y + w_item.h / 2 == pytest.approx(s_item.y + s_item.h / 2, abs=1)   # centred on each other
    assert Look.from_dict(p.look.to_dict()).logo_balance == 1.0


def test_logo_row_shrinks_to_fit_frame(tmp_path):
    from crawl.project import crawl_from_rows
    wide = _art(tmp_path / "wide.png", (2000, 100), (0, 0, 2000, 100))     # 20:1
    p = Project(rows=[["[Logo: A] | [Logo: B] | [Logo: C]", "", ""]], logos={k: wide for k in "ABC"})
    p.look.logo_balance = 0.0
    L = layout_mod.layout(crawl_from_rows(p, p.rows), p.look)
    left, right = L.logos[0].x, L.logos[-1].x + L.logos[-1].w
    assert right - left <= p.look.width * layout_mod.LOGO_ROW_MAX + 2
    assert abs((left + right) / 2 - p.look.center_x) <= 2
    assert len({(i.w, i.h) for i in L.logos}) == 1          # shrunk together, still equal


# ---- verifying a rendered file ---------------------------------------------

@pytest.mark.skipif(export.find_ffmpeg() is None, reason="ffmpeg not installed")
def test_verify_measures_uniform_motion_and_catches_judder(tmp_path):
    import subprocess
    from crawl.verify import verify
    sheet = ingest.read_rows(str(SAMPLE))[:40]
    p = Project(rows=[list(r) for r in sheet], output_dir=str(tmp_path), codec="h264")
    p.set_resolution(960, 540)
    p.timing.mode = "manual"
    p.timing.speed = 12
    b = prepare(p)
    movie = render(p, b)[0]
    good = verify(str(movie), frames=120)
    assert good.clean and good.speed == 12 and good.shifts == {12: good.frames_checked}
    assert good.frames_checked > 100 and "CLEAN" in good.summary()
    # Conform it to 29.97 by repeating frames: the same picture now moves 12, 12, 12, 0, …
    bad_path = tmp_path / "judder.mp4"
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-i", str(movie), "-vf", "fps=30000/1001",
                    "-c:v", "libx264", "-preset", "ultrafast", str(bad_path)], check=True)
    bad = verify(str(bad_path), frames=120)
    assert not bad.uniform_motion and bad.shifts.get(0, 0) > 10 and bad.speed == 12
    assert "UNEVEN MOTION" in bad.summary()


# ---- resolution ------------------------------------------------------------

def test_resize_scales_the_look_by_height_and_keeps_the_column_placed():
    look = Look()
    look.center_x = 1820                                   # 100 px left of centre at UHD
    assert look.resize(1920, 1080) == 0.5
    assert (look.width, look.height, look.center_x) == (1920, 1080, 910)     # 50 px left of centre
    assert (look.name.size, look.name.leading, look.header.tracking) == (28, 40, 2.0)
    assert (look.gutter, look.group_gap, look.logo_height, look.list_width) == (60, 60, 150, 1440)
    assert look.logo_balance == 0.85 and look.supersample == 1               # not lengths: untouched
    # Same height, wider frame: nothing scales, the column moves with the centre.
    look = Look()
    assert look.resize(4096, 2160) == 1.0
    assert (look.center_x, look.name.size, look.list_width) == (2048, 56, 2880)
    # Canvas only: sizes stay, the column is re-centred.
    look = Look()
    assert look.resize(1920, 1080, scale_look=False) == 1.0
    assert (look.center_x, look.name.size) == (960, 56)


def test_set_resolution_gives_the_same_picture_at_half_size():
    sheet = [list(r) for r in ingest.read_rows(str(SAMPLE))[:60]] + [["", "", ""], ["[Logo: A]", "", ""]]
    p = Project(rows=sheet, logo_settings={"A": {"dx": 40, "dy": -20}},
                arrangements=[])
    p.timing.mode = "manual"
    p.timing.speed = 8
    uhd = prepare(p)
    assert p.set_resolution(1920, 1080) == 0.5
    assert p.timing.speed == 4 and p.logo_settings["A"] == {"dx": 20, "dy": -10}
    hd = prepare(p)
    # Half the pixels, the same runtime (to a frame or two of rounding) and the same layout.
    assert abs(hd.solution.total_frames - uhd.solution.total_frames) <= 3
    assert abs(hd.layout.content_height * 2 - uhd.layout.content_height) <= len(uhd.layout.lines)
    assert len(hd.layout.lines) == len(uhd.layout.lines)
    assert all(abs(a.x * 2 - b.x) <= 2 for a, b in zip(hd.layout.lines, uhd.layout.lines))


# ---- logo levels -----------------------------------------------------------

def test_levels_remove_a_not_quite_black_background(tmp_path):
    import numpy as np
    from crawl import logos
    from crawl.project import crawl_from_rows
    from crawl.render import render_panel
    # A light mark on a dark-grey box with generous margin: opaque, so the box is "artwork".
    boxed = _art(tmp_path / "boxed.png", (800, 400), (200, 150, 600, 250), colour=(230, 230, 230, 255),
                 bg=(46, 46, 46, 255))
    assert logos.ink_box(boxed) == (0, 0, 800, 400)                    # untrimmed: the box counts
    level = logos.auto_black(boxed)
    assert 46 / 255 < level < 60 / 255                                 # just above the box
    assert logos.ink_box(boxed, black=level) == (200, 150, 600, 250)   # the mark alone, once the box is crushed
    assert logos.auto_black(_art(tmp_path / "clear.png", (100, 50), (10, 10, 90, 40))) == 0.0   # transparent
    assert logos.auto_black(_art(tmp_path / "white.png", (100, 50), (10, 10, 90, 40), bg=(255, 255, 255, 255))) == 0.0
    # …but a white background is removable once inverted.
    paper = _art(tmp_path / "paper.png", (100, 50), (10, 10, 90, 40), colour=(20, 20, 20, 255), bg=(244, 244, 244, 255))
    assert 0 < logos.auto_black(paper, invert=True) < 0.1
    assert logos.ink_box(paper, invert=True, black=logos.auto_black(paper, invert=True)) == (10, 10, 90, 40)

    p = Project(rows=[["[Logo: A]", "", ""]], logos={"A": boxed})
    look = p.look

    def drawn():
        L = layout_mod.layout(crawl_from_rows(p, p.rows), look)
        item = L.logos[0]
        img = np.asarray(render_panel(L, look, item.y - 20, item.y + item.h + 20, 1.0, 1, colour=False))
        return item, img

    item, img = drawn()
    assert img[25, item.x + 5] > 30                                     # the grey box shows
    p.logo_settings = {"A": {"black": int(round(level * 100)), "white": 90}}
    item2, img2 = drawn()
    assert img2.max() == 255                                            # white point lifted 230 → full
    assert img2[:, : item2.x - 2].max() == 0 and img2[:18].max() == 0    # nothing but the mark remains
    assert item2.w / item2.h == pytest.approx(4.0, abs=0.05) and item.w / item.h == pytest.approx(2.0, abs=0.05)


# ---- portable projects -----------------------------------------------------

def test_project_paths_are_relative_inside_its_folder_and_survive_a_move(tmp_path):
    import json, shutil
    job = tmp_path / "job"
    (job / "logos").mkdir(parents=True)
    _logo_png(job / "logos" / "studio.png", (240, 240, 240, 255))
    outside = _logo_png(tmp_path / "elsewhere.png", (240, 240, 240, 255))
    (job / "credits.csv").write_text(",Lead,,Ada\n,[Logo: A] | [Logo: B],,\n")
    p = Project(source=str(job / "credits.csv"), output_dir=str(job / "render"),
                logos={"A": str(job / "logos" / "studio.png"), "B": outside})
    path = job / "show.crawl.json"
    p.save(str(path))
    saved = json.loads(path.read_text())
    assert saved["source"] == "credits.csv" and saved["output_dir"] == "render"
    assert saved["logos"] == {"A": "logos/studio.png", "B": outside}        # outside the folder: absolute
    assert Project.load(str(path)).logos["A"] == str(job / "logos" / "studio.png")

    # Collect brings the outside file in; then the whole folder can move.
    assert p.collect(str(path)) == ["elsewhere.png"]
    assert p.logos["B"] == str(job / "show_assets" / "elsewhere.png") and p.collect(str(path)) == []
    p.save(str(path))
    assert json.loads(path.read_text())["logos"]["B"] == "show_assets/elsewhere.png"
    moved = tmp_path / "another machine" / "job"
    shutil.copytree(job, moved)
    shutil.rmtree(job)
    q = Project.load(str(moved / "show.crawl.json"))
    assert q.load_notes == []
    assert q.source == str(moved / "credits.csv") and q.output_dir == str(moved / "render")
    assert all(Path(f).is_file() and str(moved) in f for f in q.logos.values())
    assert [lg.path for lg in __import__("crawl.project", fromlist=["x"]).load_crawl(q).logos()] == list(q.logos.values())


def test_loading_repairs_moved_files_and_reports_missing_ones(tmp_path):
    from crawl.project import findings_for, load_crawl
    old = tmp_path / "old"
    old.mkdir()
    p = Project(source=str(old / "credits.csv"), output_dir="/Volumes/NotHere/render",
                logos={"A": "/Users/someone-else/art/studio.png", "B": "/Users/someone-else/art/gone.png"},
                rows=None, base_rows=[["Lead", "", "Ada"], ["[Logo: A] | [Logo: B]", "", ""]])
    path = tmp_path / "show.crawl.json"
    p.save(str(path))
    _logo_png(tmp_path / "studio.png", (240, 240, 240, 255))            # same name, now beside the project
    q = Project.load(str(path))
    assert q.logos["A"] == str(tmp_path / "studio.png")                  # relinked by file name
    assert q.rows == [["Lead", "", "Ada"], ["[Logo: A] | [Logo: B]", "", ""]]   # sheet gone: credits from the project
    assert q.output_dir == str(tmp_path / "render")
    notes = " ".join(q.load_notes)
    assert "Found 1 moved logo" in notes and "not found for: B" in notes and "credits.csv was not found" in notes
    found = findings_for(q, load_crawl(q))
    assert [(f.severity, f.key) for f in found] == [("error", "logo-missing|b")]
    assert "load_notes" not in q.to_dict()


# ---- cancelling an export --------------------------------------------------

@pytest.mark.skipif(export.find_ffmpeg() is None, reason="ffmpeg not installed")
def test_cancel_stops_the_encoder_and_removes_the_partial_file(tmp_path):
    sheet = [list(r) for r in ingest.read_rows(str(SAMPLE))[:40]]
    p = Project(rows=sheet, output_dir=str(tmp_path), codec="h264", basename="show")
    p.set_resolution(960, 540)
    p.timing.mode = "manual"
    p.timing.speed = 4
    b = prepare(p)
    assert b.solution.total_frames > 200
    seen = []

    def progress(stage, i, n):
        seen.append((stage, i))

    polls = {"n": 0}

    def cancel():                       # let it encode a while, then stop
        polls["n"] += 1
        return polls["n"] > 120

    with pytest.raises(export.ExportCancelled):
        render(p, b, progress, cancel)
    assert any(stage == "encode" for stage, _ in seen)                 # it had really started
    assert max(i for stage, i in seen if stage == "encode") < b.solution.total_frames
    assert list(tmp_path.iterdir()) == []                               # no movie, no report
    # Cancelling before drawing starts leaves nothing either, and a PNG folder it created is removed.
    p.codec = "png"
    with pytest.raises(export.ExportCancelled):
        render(p, b, None, lambda: True)
    assert list(tmp_path.iterdir()) == []
    # Not cancelling still works.
    p.codec = "h264"
    paths = render(p, b, None, lambda: False)
    assert (tmp_path / "show.mp4").is_file() and len(paths) == 1


def test_default_font_falls_back_but_named_fonts_do_not(monkeypatch):
    from crawl import fonts, text
    real_find = fonts.find

    def no_helvetica_neue(family, style="Regular"):
        return None if family.lower() == "helvetica neue" else real_find(family, style)

    monkeypatch.setattr(fonts, "find", no_helvetica_neue)
    text._load.cache_clear()
    face = fonts.resolve_default("Helvetica Neue:Bold")
    assert face is not None and face.family != "Helvetica Neue"
    assert text.font_for(Look().header) is not None               # the default look still works
    with pytest.raises(text.FontError):
        text._load("Gotham Nonexistent:Black", 40)                 # a named font does not quietly change
    text._load.cache_clear()
